"""Real in-memory fixtures and extraction admission/resource boundaries, without a model."""
import base64
import io
import json
import subprocess
import struct
import unittest
import zipfile
from unittest import mock

from serve import file_extract as extraction


def request(name, raw):
    return {'name': name, 'data': base64.b64encode(raw).decode('ascii')}


def pdf(text=None, pages=1):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=200, height=200)
        if text:
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                                     NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
            stream = DecodedStreamObject()
            stream.set_data(f'BT /F1 12 Tf 10 100 Td ({text}) Tj ET'.encode())
            page[NameObject('/Contents')] = stream
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class Extraction(unittest.TestCase):
    def test_text_and_code_utf8_utf16_and_cap(self):
        for name, raw in [('a.py', '中文\nprint(1)'.encode()), ('a.txt', '中文'.encode('utf-16'))]:
            result = extraction.extract(request(name, raw))
            self.assertIn('中文', result['text'])
            self.assertFalse(result['truncated'])
        result = extraction._extract_local(request('large.txt', ('中' * 200000).encode()))
        self.assertLessEqual(len(result['text'].encode()), extraction.MAX_TEXT)
        self.assertTrue(result['truncated'])
        self.assertTrue(result['warnings'])

    def test_strict_input_and_magic(self):
        invalid = [({}, 400), ({'name': '../a.txt', 'data': 'YWJj'}, 400),
                   ({'name': 'a.txt', 'data': 'data:text/plain;base64,YWJj'}, 400),
                   ({'name': 'a.txt', 'data': 'YWJj', 'path': 'C:/secret'}, 400),
                   (request('a.exe', b'abc'), 415), (request('a.txt', b''), 422)]
        for req, status in invalid:
            with self.subTest(req=req), self.assertRaises(extraction.ExtractionError) as error:
                extraction.validate_request(req)
            self.assertEqual(error.exception.status, status)
        for name, raw in [('a.pdf', b'abc'), ('a.docx', b'abc'), ('a.xlsx', b'abc'), ('a.txt', b'%PDF-1.7'), ('a.txt', b'\x00abc')]:
            with self.assertRaises(extraction.ExtractionError):
                extraction._extract_local(request(name, raw))

    def test_docx_chinese_table_in_document_order(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('前文 中文')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '甲'
        table.cell(0, 1).text = '乙'
        doc.add_paragraph('后文')
        buffer = io.BytesIO()
        doc.save(buffer)
        result = extraction.extract(request('table.docx', buffer.getvalue()))
        self.assertLess(result['text'].index('前文'), result['text'].index('甲\t乙'))
        self.assertLess(result['text'].index('甲\t乙'), result['text'].index('后文'))

    def test_xlsx_multiple_sheets_and_no_formula_execution(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = '中文'
        book.active.append(['姓名', '张三', 42])
        second = book.create_sheet('Second')
        second.append(['safe', '=1+1'])
        buffer = io.BytesIO()
        book.save(buffer)
        result = extraction.extract(request('sheets.xlsx', buffer.getvalue()))
        self.assertIn('[Sheet: 中文]', result['text'])
        self.assertIn('张三\t42', result['text'])
        self.assertIn('[Sheet: Second]', result['text'])
        self.assertNotIn('=1+1', result['text'])
        self.assertIn('not executed', result['warnings'][0])

    def test_pdf_text_and_scanned_fallback_actual_pixels(self):
        result = extraction.extract(request('text.pdf', pdf('HELLO-7421')))
        self.assertIn('HELLO-7421', result['text'])
        self.assertNotIn('images', result)
        result = extraction.extract(request('scan.pdf', pdf(pages=9)))
        self.assertEqual(len(result['images']), 8)
        self.assertTrue(result['truncated'])
        from PIL import Image
        total = 0
        for image in result['images']:
            encoded = image['url'].split(',', 1)[1]
            total += len(encoded)
            with Image.open(io.BytesIO(base64.b64decode(encoded))) as pixels:
                self.assertLessEqual(max(pixels.size), 1600)
        self.assertLessEqual(total, extraction.MAX_IMAGES)

    def test_encrypted_and_malformed_pdf_safe_errors(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.encrypt('fixture')
        buffer = io.BytesIO()
        writer.write(buffer)
        for raw in [buffer.getvalue(), b'%PDF-1.7\nmalformed']:
            with self.assertRaises(extraction.ExtractionError) as error:
                extraction.extract(request('a.pdf', raw))
            self.assertEqual(error.exception.status, 422)

    def test_office_zip_traversal_xxe_bomb_without_disk_extraction(self):
        for extra, payload in [('../escape.txt', b'x'), ('word/evil.xml', b'<!DOCTYPE a [<!ENTITY e SYSTEM "file:///secret">]><a>&e;</a>'),
                               ('word/bomb.xml', b' ' * 1000000)]:
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr('[Content_Types].xml', '<Types/>')
                archive.writestr('word/document.xml', '<document/>')
                archive.writestr(extra, payload)
            with self.assertRaises(extraction.ExtractionError) as error:
                extraction.extract(request('bad.docx', buffer.getvalue()))
            self.assertIn(error.exception.status, (413, 422))

    def test_encrypted_zip_flag_and_decoded_input_limit(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('[Content_Types].xml', '<Types/>')
            archive.writestr('word/document.xml', '<document/>')
        raw = bytearray(buffer.getvalue())
        central = raw.index(b'PK\x01\x02')
        flags = struct.unpack_from('<H', raw, central + 8)[0]
        struct.pack_into('<H', raw, central + 8, flags | 1)
        with self.assertRaises(extraction.ExtractionError) as error:
            extraction.extract(request('encrypted.docx', raw))
        self.assertEqual(error.exception.status, 422)
        with mock.patch.object(extraction, 'MAX_INPUT', 2):
            with self.assertRaises(extraction.ExtractionError) as error:
                extraction.validate_request(request('a.txt', b'abc'))
        self.assertEqual(error.exception.status, 413)

    def test_missing_parser_dependency_is_503(self):
        fixture = request('a.pdf', pdf())
        output = io.StringIO()
        stdin = mock.Mock(buffer=io.BytesIO(json.dumps(fixture).encode()))
        with mock.patch.dict('sys.modules', {'pypdf': None}), mock.patch.object(extraction, '_limit_worker'), \
                mock.patch.object(extraction.sys, 'stdin', stdin), mock.patch.object(extraction.sys, 'stdout', output):
            extraction._worker()
        self.assertEqual(json.loads(output.getvalue())['status'], 503)

    def test_memory_only_parsers_do_not_allocate_upload_tempfiles(self):
        fixture = request('a.pdf', pdf('MEMORY-ONLY'))
        with mock.patch('tempfile.NamedTemporaryFile', side_effect=AssertionError('upload persisted')), \
                mock.patch('tempfile.TemporaryDirectory', side_effect=AssertionError('upload persisted')):
            self.assertIn('MEMORY-ONLY', extraction._extract_local(fixture)['text'])

    def test_timeout_kills_worker_and_concurrency_rejects(self):
        child = mock.MagicMock()
        child.__enter__.return_value = child
        child.communicate.side_effect = [subprocess.TimeoutExpired('worker', 30), (b'', None)]
        with mock.patch.object(extraction.subprocess, 'Popen', return_value=child):
            with self.assertRaises(extraction.ExtractionError) as error:
                extraction.extract(request('a.txt', b'hello'))
            self.assertEqual(error.exception.status, 422)
            child.kill.assert_called_once()
        extraction._slots.acquire()
        extraction._slots.acquire()
        try:
            with self.assertRaises(extraction.ExtractionError) as error:
                extraction.extract(request('a.txt', b'hello'))
            self.assertEqual(error.exception.status, 503)
        finally:
            extraction._slots.release()
            extraction._slots.release()


if __name__ == '__main__':
    unittest.main()
