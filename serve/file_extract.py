"""Bounded local document extraction; uploaded bytes live only in process memory."""
from __future__ import annotations

import base64
import binascii
import io
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import threading
import zipfile

from serve.winjob import contain

MAX_INPUT = 20 * 1024 * 1024
MAX_BODY = 28 * 1024 * 1024
MAX_TEXT = 512 * 1024
MAX_IMAGES = 8 * 1024 * 1024
TIMEOUT = 30
_slots = threading.BoundedSemaphore(2)
TEXT_EXTENSIONS = {'.txt', '.md', '.py', '.js', '.jsx', '.ts', '.tsx', '.json', '.csv', '.yaml', '.yml',
                   '.toml', '.ini', '.html', '.css', '.sql', '.sh', '.ps1', '.c', '.cpp', '.h', '.hpp',
                   '.rs', '.go', '.java', '.log', '.xml', '.svg', '.ipynb', '.rb', '.php', '.bat'}


class ExtractionError(ValueError):
    def __init__(self, status, message):
        super().__init__(message)
        self.status = status


def validate_request(req):
    if not isinstance(req, dict) or set(req) != {'name', 'data'}:
        raise ExtractionError(400, 'send only name and base64 data')
    name, data = req['name'], req['data']
    if (not isinstance(name, str) or not name or len(name) > 255 or
            any(c in name for c in '/\\:') or any(ord(c) < 32 for c in name) or name in {'.', '..'}):
        raise ExtractionError(400, 'send a filename without a path')
    if not isinstance(data, str):
        raise ExtractionError(400, 'data must be base64 text')
    if len(data) > 4 * math.ceil(MAX_INPUT / 3):
        raise ExtractionError(413, 'file exceeds the 20 MiB limit')
    try:
        raw = base64.b64decode(data, validate=True)
    except (ValueError, binascii.Error):
        raise ExtractionError(400, 'data must contain base64 only') from None
    if not raw:
        raise ExtractionError(422, 'file is empty')
    if len(raw) > MAX_INPUT:
        raise ExtractionError(413, 'file exceeds the 20 MiB limit')
    suffix = Path(name).suffix.lower()
    if suffix not in TEXT_EXTENSIONS | {'.pdf', '.docx', '.xlsx'} and name != '.env':
        raise ExtractionError(415, 'supported files: PDF, DOCX, XLSX, text and source code')
    return name, raw, suffix


def extract(req):
    if not _slots.acquire(blocking=False):
        raise ExtractionError(503, 'two files are already being extracted; try again shortly')
    try:
        validate_request(req)
        flags = subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0
        try:
            with subprocess.Popen([sys.executable, '-B', '-m', 'serve.file_extract', '--worker'],
                                  cwd=Path(__file__).resolve().parents[1], stdin=subprocess.PIPE,
                                  stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                                  creationflags=flags, env={**os.environ, 'PYTHONUTF8': '1'}) as child:
                contain(child)  # existing lifetime owner; child also installs its own memory-limited nested job
                try:
                    output, _ = child.communicate(json.dumps(req).encode('utf-8'), timeout=TIMEOUT)
                except subprocess.TimeoutExpired:
                    child.kill()
                    child.communicate()
                    raise ExtractionError(422, 'file extraction exceeded the 30 second limit') from None
                if child.returncode or len(output) > MAX_IMAGES + 2 * MAX_TEXT + 1024 * 1024:
                    raise ExtractionError(422, 'file extraction stopped at its resource limit')
                response = json.loads(output)
        except (OSError, json.JSONDecodeError):
            raise ExtractionError(503, 'local extraction worker is unavailable') from None
        if 'error' in response:
            raise ExtractionError(response['status'], response['error'])
        return response
    finally:
        _slots.release()


def _limit_worker():
    # Applied before importing parsers. A separate child owns its memory-limited job.
    if os.name == 'nt':
        import ctypes
        from serve import winjob
        job = winjob._make_job()
        if not job:
            raise ExtractionError(503, 'cannot establish the extraction memory limit')
        info = winjob._ExtendedLimits()
        info.BasicLimitInformation.LimitFlags = 0x2000 | 0x100  # kill on close; process committed memory
        info.ProcessMemoryLimit = 768 * 1024 * 1024
        if not winjob._k32.SetInformationJobObject(job, 9, ctypes.byref(info), ctypes.sizeof(info)):
            raise ExtractionError(503, 'cannot establish the extraction memory limit')
        winjob._k32.GetCurrentProcess.restype = ctypes.c_void_p
        if not winjob._k32.AssignProcessToJobObject(job, winjob._k32.GetCurrentProcess()):
            raise ExtractionError(503, 'cannot establish the extraction memory limit')
        globals()['_worker_job'] = job  # retain handle until child exit
    else:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024,) * 2)


def _office_zip(raw, required):
    from defusedxml.ElementTree import fromstring
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw))
        entries = archive.infolist()
        if len(entries) > 2000 or sum(e.file_size for e in entries) > 40 * 1024 * 1024:
            raise ExtractionError(413, 'Office file exceeds the expanded size or entry limit')
        names = set()
        for entry in entries:
            name = entry.filename
            if (name in names or name.startswith(('/', '\\')) or '\\' in name or ':' in name or
                    '..' in name.split('/') or entry.flag_bits & 1):
                raise ExtractionError(422, 'Office ZIP contains unsafe or encrypted entries')
            names.add(name)
            if entry.file_size > 16 * 1024 * 1024 or entry.file_size > max(1, entry.compress_size) * 200:
                raise ExtractionError(413, 'Office file exceeds the per-entry expansion limit')
            if name.lower().endswith(('.xml', '.rels')):
                xml = archive.read(entry)
                fromstring(xml, forbid_dtd=True, forbid_entities=True, forbid_external=True)
        if required not in names or '[Content_Types].xml' not in names:
            raise ExtractionError(422, 'filename and Office file format do not match')
    finally:
        if 'archive' in locals():
            archive.close()


def _extract_local(req):
    name, raw, suffix = validate_request(req)
    warnings, images, chunks = [], [], []
    truncated = False
    used = 0

    def add(text):
        nonlocal used, truncated
        encoded = text.encode('utf-8')
        remaining = MAX_TEXT - used
        if len(encoded) > remaining:
            text = encoded[:remaining].decode('utf-8', errors='ignore')
            truncated = True
        chunks.append(text)
        used += len(text.encode('utf-8'))
        return not truncated

    if suffix == '.pdf':
        if not raw.startswith(b'%PDF-'):
            raise ExtractionError(422, 'filename and PDF file format do not match')
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw), strict=False)
        if reader.is_encrypted:
            raise ExtractionError(422, 'encrypted PDFs are not supported; provide a decrypted copy')
        count = len(reader.pages)
        stream_bytes = 0
        scanned = []
        for index, page in enumerate(reader.pages):
            if index >= 200:
                warnings.append('Only the first 200 PDF pages were inspected.')
                truncated = True
                break
            contents = page.get_contents()
            if contents:
                size = len(contents.get_data())
                stream_bytes += size
                if size > 8 * 1024 * 1024 or stream_bytes > 40 * 1024 * 1024:
                    raise ExtractionError(413, 'PDF exceeds the decoded content stream limit')
            text = page.extract_text() or ''
            if text.strip():
                if not add(f'\n[Page {index + 1}]\n{text}\n'):
                    break
            else:
                scanned.append(index)
        if scanned:
            import pypdfium2 as pdfium
            document = pdfium.PdfDocument(raw)
            encoded_total = 0
            try:
                for index in scanned[:8]:
                    page = document[index]
                    bitmap = None
                    try:
                        width, height = page.get_size()
                        if not all(math.isfinite(n) and n > 0 for n in (width, height)):
                            raise ExtractionError(422, 'PDF page has invalid dimensions')
                        bitmap = page.render(scale=min(2.0, 1599 / max(width, height)))
                        image = bitmap.to_pil()
                        try:
                            if max(image.size) > 1600:
                                raise ExtractionError(422, 'PDF render exceeds the dimension limit')
                            buffer = io.BytesIO()
                            image.save(buffer, format='PNG')
                            encoded = base64.b64encode(buffer.getvalue()).decode('ascii')
                        finally:
                            image.close()
                        if encoded_total + len(encoded) > MAX_IMAGES:
                            truncated = True
                            warnings.append('Scanned PDF images reached the 8 MiB limit.')
                            break
                        encoded_total += len(encoded)
                        images.append({'name': f'{name} · page {index + 1}', 'url': 'data:image/png;base64,' + encoded})
                        add(f'\n[Page {index + 1}: image attached; no extractable text]\n')
                    finally:
                        if bitmap is not None:
                            bitmap.close()
                        page.close()
            finally:
                document.close()
            warnings.append('Pages without extractable text were attached as images; reading them requires vision.')
            if len(scanned) > 8:
                truncated = True
                warnings.append('Only the first 8 pages without text were rendered.')
        if count == 0:
            raise ExtractionError(422, 'PDF contains no pages')
    elif suffix == '.docx':
        if not raw.startswith(b'PK'):
            raise ExtractionError(422, 'filename and DOCX file format do not match')
        _office_zip(raw, 'word/document.xml')
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(io.BytesIO(raw))
        for child in doc.element.body:
            if child.tag.endswith('}p'):
                if not add(Paragraph(child, doc).text + '\n'):
                    break
            elif child.tag.endswith('}tbl'):
                for row in Table(child, doc).rows:
                    if not add('\t'.join(cell.text for cell in row.cells) + '\n'):
                        break
            if truncated:
                break
        warnings.append('DOCX extraction includes body paragraphs and tables; embedded images are not extracted.')
    elif suffix == '.xlsx':
        if not raw.startswith(b'PK'):
            raise ExtractionError(422, 'filename and XLSX file format do not match')
        _office_zip(raw, 'xl/workbook.xml')
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        cells = 0
        try:
            for sheet_index, sheet in enumerate(book.worksheets):
                if sheet_index >= 50:
                    truncated = True
                    warnings.append('Only the first 50 worksheets were extracted.')
                    break
                add(f'\n[Sheet: {sheet.title}]\n')
                for row_index, row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row or 0, 10000),
                                                               max_col=min(sheet.max_column or 0, 200), values_only=True)):
                    cells += len(row)
                    if cells > 100000:
                        truncated = True
                        break
                    if not add('\t'.join('' if value is None else str(value) for value in row) + '\n'):
                        break
                if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 200:
                    truncated = True
                if truncated:
                    break
        finally:
            book.close()
        warnings.append('Spreadsheet formulas are not executed; only stored cached values are read.')
        if cells > 100000:
            warnings.append('Spreadsheet extraction reached the 100,000 cell limit.')
    else:
        if raw.startswith((b'%PDF-', b'PK\x03\x04', b'\xd0\xcf\x11\xe0')):
            raise ExtractionError(422, 'binary document does not match the text filename')
        try:
            encoding = 'utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
            text = raw.decode(encoding)
        except UnicodeError:
            raise ExtractionError(422, 'text files must use UTF-8 or UTF-16') from None
        if '\x00' in text:
            raise ExtractionError(422, 'binary content is not a text file')
        add(text)
    if truncated:
        warnings.append('Extraction was limited; some content is omitted.')
    return {'name': name, 'text': ''.join(chunks), 'warnings': warnings, 'truncated': truncated,
            'format': suffix.lstrip('.') or 'text', **({'images': images} if images else {})}


def _worker():
    try:
        _limit_worker()
        req = json.loads(sys.stdin.buffer.read(MAX_BODY + 1))
        result = _extract_local(req)
    except ExtractionError as error:
        result = {'error': str(error), 'status': error.status}
    except ImportError:
        result = {'error': 'local document extraction dependencies are not installed', 'status': 503}
    except Exception:
        result = {'error': 'file is malformed or cannot be extracted within the local limits', 'status': 422}
    sys.stdout.write(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__' and sys.argv[1:] == ['--worker']:
    _worker()
